import docx
from copy import deepcopy
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls

doc = docx.Document(r'D:\Dự án bảo tàng\Bao_cao_Mo_phong_Bao_tang_Trang_phuc_Tay_Nguyen_Chuong_3.docx')

# Let's test cloning paragraph and inserting table
last_p = doc.paragraphs[-1]
print('Last paragraph text:', last_p.text[:50])

# Clone paragraph
p_elem = deepcopy(last_p._p)
last_p._p.addnext(p_elem)
new_p = doc.paragraphs[[x._p for x in doc.paragraphs].index(p_elem)]
for child in list(new_p._p):
    if child.tag == qn('w:r') or child.tag == qn('w:hyperlink'):
        new_p._p.remove(child)
r = new_p.add_run("Test after last paragraph")

# Create a small table
tbl = doc.add_table(rows=2, cols=2)
tbl.rows[0].cells[0].text = "Col 1"
tbl.rows[0].cells[1].text = "Col 2"
tbl.rows[1].cells[0].text = "Val 1"
tbl.rows[1].cells[1].text = "Val 2"

# Move table after new_p
new_p._p.addnext(tbl._tbl)

# Add another paragraph after table._tbl
p_after = deepcopy(last_p._p)
tbl._tbl.addnext(p_after)
new_p2 = doc.paragraphs[[x._p for x in doc.paragraphs].index(p_after)]
for child in list(new_p2._p):
    if child.tag == qn('w:r') or child.tag == qn('w:hyperlink'):
        new_p2._p.remove(child)
r2 = new_p2.add_run("Paragraph after table")

doc.save(r'D:\Dự án bảo tàng\Assets\_docx_qa\test_out.docx')
print("Successfully generated test_out.docx!")
