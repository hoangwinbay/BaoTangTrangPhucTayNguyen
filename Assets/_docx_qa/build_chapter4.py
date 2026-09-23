import os
from copy import deepcopy
from docx import Document
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Inches, RGBColor

SOURCE = r"D:\Dự án bảo tàng\Bao_cao_Mo_phong_Bao_tang_Trang_phuc_Tay_Nguyen_Chuong_3.docx"
OUTPUT_QA = r"D:\Dự án bảo tàng\Assets\_docx_qa\Bao_cao_Mo_phong_Bao_tang_Trang_phuc_Tay_Nguyen_Chuong_4.docx"
OUTPUT_ROOT = r"D:\Dự án bảo tàng\Bao_cao_Mo_phong_Bao_tang_Trang_phuc_Tay_Nguyen_Chuong_4.docx"

doc = Document(SOURCE)

def clone_after(anchor, template, text, style=None, page_break=False):
    """Clone an existing paragraph to preserve font, size, line spacing, margins, indents."""
    p = deepcopy(template._p)
    anchor._p.addnext(p)
    new = doc.paragraphs[[x._p for x in doc.paragraphs].index(p)]
    for child in list(new._p):
        if child.tag == qn('w:r') or child.tag == qn('w:hyperlink'):
            new._p.remove(child)
    if style:
        new.style = style
        
    rPr_copy = deepcopy(template.runs[0]._element.rPr) if template.runs and template.runs[0]._element.rPr is not None else None
    
    if '\t' in text:
        parts = text.split('\t')
        title_part = parts[0].strip()
        page_part = parts[-1].strip()
        
        r1 = new.add_run(title_part)
        if rPr_copy is not None:
            r1._element.get_or_add_rPr().append(deepcopy(rPr_copy))
            
        r2 = new.add_run()
        if rPr_copy is not None:
            r2._element.get_or_add_rPr().append(deepcopy(rPr_copy))
        r2.add_tab()
        
        r3 = new.add_run(page_part)
        if rPr_copy is not None:
            r3._element.get_or_add_rPr().append(deepcopy(rPr_copy))
    else:
        r = new.add_run()
        if rPr_copy is not None:
            r._element.get_or_add_rPr().append(deepcopy(rPr_copy))
        if page_break:
            r.add_break(WD_BREAK.PAGE)
        r.add_text(text)
    return new

def insert_styled_table(anchor_elem, headers, data, col_widths=None):
    """Creates a concise styled table matching academic layout."""
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    tblPr = table._tbl.tblPr
    borders_xml = parse_xml(
        r'<w:tblBorders %s>'
        r'<w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(borders_xml)
    
    tblCellMar = parse_xml(
        r'<w:tblCellMar %s>'
        r'<w:top w:w="120" w:type="dxa"/>'
        r'<w:bottom w:w="120" w:type="dxa"/>'
        r'<w:left w:w="160" w:type="dxa"/>'
        r'<w:right w:w="160" w:type="dxa"/>'
        r'</w:tblCellMar>' % nsdecls('w')
    )
    tblPr.append(tblCellMar)
    
    hdr_row = table.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
    
    for c_idx, title in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.text = title
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(r'<w:shd %s w:fill="D9D9D9"/>' % nsdecls('w')))
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if p.runs:
            p.runs[0].font.name = "Times New Roman"
            p.runs[0].font.size = Pt(12)
            p.runs[0].font.bold = True
            
    for r_idx, row_vals in enumerate(data):
        row = table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.text = str(val)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            str_val = str(val).strip()
            if str_val in ["Đạt", "Chờ đo", "Pass", "P0", "P1"] or len(str_val) <= 4 or str_val.startswith("TC-"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if p.runs:
                p.runs[0].font.name = "Times New Roman"
                p.runs[0].font.size = Pt(12)
                if str_val in ["Đạt", "Pass"]:
                    p.runs[0].font.bold = True
                    
    if col_widths:
        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                if c_idx < len(row.cells):
                    row.cells[c_idx].width = Inches(w)
                    
    anchor_elem.addnext(table._tbl)
    return table

# Templates
toc_chapter = doc.paragraphs[52]
toc_l2 = doc.paragraphs[55]
toc_l3 = doc.paragraphs[56]

chapter = doc.paragraphs[187]   # Chapter title style (16pt, Bold)
h2 = doc.paragraphs[188]        # Heading 2 style (14pt, Bold)
h3 = doc.paragraphs[189]        # Heading 3 style (13pt, Bold)
body = doc.paragraphs[190]      # Body style (13pt, firstLine 720 dxa)
table_caption = doc.paragraphs[185] # Table caption style (Bold, Center)

# ---------------------------------------------------------
# 1. Table of Contents
# ---------------------------------------------------------
toc_anchor = doc.paragraphs[90]
toc_entries = [
    ('CHƯƠNG 4: KIỂM THỬ VÀ ĐÁNH GIÁ ỨNG DỤNG VR\t28', toc_chapter),
    ('4.1. Xây dựng kế hoạch kiểm thử\t28', toc_l2),
    ('4.1.1. Mục tiêu và phạm vi kiểm thử\t28', toc_l3),
    ('4.1.2. Môi trường và thiết bị kiểm thử\t28', toc_l3),
    ('4.1.3. Phương pháp và tiêu chí kiểm thử\t28', toc_l3),
    ('4.2. Kiểm thử chức năng\t28', toc_l2),
    ('4.2.1. Kiểm thử khởi động và điều khiển\t28', toc_l3),
    ('4.2.2. Kiểm thử di chuyển và tương tác VR\t29', toc_l3),
    ('4.2.3. Kiểm thử Spatial UI và các chức năng chính\t30', toc_l3),
    ('4.2.4. Kiểm thử âm thanh và phản hồi xúc giác\t31', toc_l3),
    ('4.3. Kiểm thử hiệu năng\t31', toc_l2),
    ('4.3.1. Kiểm thử FPS và Frame Time\t31', toc_l3),
    ('4.3.2. Kiểm thử độ trễ và khả năng phản hồi\t32', toc_l3),
    ('4.3.3. Kiểm thử CPU, GPU và bộ nhớ\t33', toc_l3),
    ('4.4. Đánh giá trải nghiệm người dùng\t33', toc_l2),
    ('4.4.1. Đánh giá tính dễ sử dụng\t33', toc_l3),
    ('4.4.2. Đánh giá mức độ tương tác và trực quan\t33', toc_l3),
    ('4.4.3. Đánh giá mức độ thoải mái khi sử dụng VR\t33', toc_l3),
    ('4.5. Tổng hợp và đánh giá kết quả\t34', toc_l2),
    ('4.5.1. Tổng hợp kết quả kiểm thử\t34', toc_l3),
    ('4.5.2. Những kết quả đạt được\t34', toc_l3),
    ('4.5.3. Hạn chế của ứng dụng\t35', toc_l3),
]
for text, template in toc_entries:
    toc_anchor = clone_after(toc_anchor, template, text)

# ---------------------------------------------------------
# 2. Build Chapter 4 Content (Concise, authentic, no fabrication)
# ---------------------------------------------------------
curr_elem = doc.paragraphs[-1]._p

def append_p(template, text, page_break=False):
    global curr_elem
    p = deepcopy(template._p)
    curr_elem.addnext(p)
    new_p = doc.paragraphs[[x._p for x in doc.paragraphs].index(p)]
    for child in list(new_p._p):
        if child.tag == qn('w:r') or child.tag == qn('w:hyperlink'):
            new_p._p.remove(child)
    r = new_p.add_run()
    if template.runs and template.runs[0]._element.rPr is not None:
        r._element.get_or_add_rPr().append(deepcopy(template.runs[0]._element.rPr))
    if page_break:
        r.add_break(WD_BREAK.PAGE)
    r.add_text(text)
    curr_elem = new_p._p
    return new_p

def append_table(caption_text, headers, data, col_widths=None):
    global curr_elem
    cap_p = append_p(table_caption, caption_text)
    tbl = insert_styled_table(cap_p._p, headers, data, col_widths)
    curr_elem = tbl._tbl
    return tbl

# ==================== NỘI DUNG CHƯƠNG 4 ====================

# Tiêu đề chương
append_p(chapter, "CHƯƠNG 4: KIỂM THỬ VÀ ĐÁNH GIÁ ỨNG DỤNG VR", page_break=True)

# 4.1. Xây dựng kế hoạch kiểm thử
append_p(h2, "4.1. Xây dựng kế hoạch kiểm thử")

append_p(h3, "4.1.1. Mục tiêu và phạm vi kiểm thử")
append_p(body, "Mục tiêu của giai đoạn kiểm thử là xác minh tính đúng đắn của các tính năng tương tác, kiểm tra độ ổn định khung hình và đánh giá trải nghiệm thực tế của người dùng khi tham quan bảo tàng ảo. Phạm vi kiểm thử tập trung vào cảnh chính (SampleScene), bao gồm: cơ chế nhận diện môi trường chạy (kính VR hoặc máy tính PC), hệ thống di chuyển theo hướng nhìn, tương tác raycast với 12 bệ trưng bày trang phục dân tộc, bảng thông tin pop-up, cửa kính tự động, nhân vật NPC và màn hình Tivi phát video tư liệu.")

append_p(h3, "4.1.2. Môi trường và thiết bị kiểm thử")
append_p(body, "Quá trình kiểm thử được tiến hành trên hai môi trường mục tiêu của dự án: (1) Kính thực tại ảo độc lập Meta Quest chạy hệ điều hành Android (thử nghiệm qua gói cài đặt APK standalone và liên kết Oculus Link); (2) Máy tính cá nhân chạy Windows với cấu hình phát triển (CPU Intel, GPU hỗ trợ DirectX 11/12, cài đặt Unity 6 phiên bản 6000.5.9f1 và URP 17.5.0). Các công cụ hỗ trợ theo dõi gồm Unity Profiler và OVR Metrics Tool. [Ghi chú: Nhóm bổ sung thông số phần cứng chi tiết của máy tính và dòng kính thực tế dùng trong đợt kiểm thử tại đây].")

append_p(h3, "4.1.3. Phương pháp và tiêu chí kiểm thử")
append_p(body, "Nhóm áp dụng phương pháp kiểm thử chức năng (kiểm tra từng ca sử dụng theo kịch bản tham quan), kiểm thử hiệu năng (đo đạc tốc độ khung hình, thời gian dựng hình, mức chiếm dụng bộ nhớ qua Profiler) và ghi nhận ý kiến người trải nghiệm. Tiêu chí nghiệm thu gồm: toàn bộ các chức năng tương tác cốt lõi vận hành ổn định không phát sinh lỗi crash; tốc độ khung hình hướng tới mức chuẩn 72 FPS trên kính Quest và tối thiểu 60 FPS trên PC; giao diện dễ hiểu và không gây cảm giác khó chịu hay mất thăng bằng khi di chuyển.")

# 4.2. Kiểm thử chức năng
append_p(h2, "4.2. Kiểm thử chức năng")

append_p(h3, "4.2.1. Kiểm thử khởi động và điều khiển")
append_p(body, "Cơ chế quản lý khởi động MuseumVRRigManager được kiểm tra về khả năng tự động nhận diện thiết bị qua VRInputHelper. Khi chạy trên kính VR, hệ thống kích hoạt XR Origin và bộ điều khiển tay cầm, đồng thời tắt FPSController của PC để tránh xung đột camera và âm thanh. Khi chạy trên PC không cắm kính, hệ thống duy trì điều khiển bằng chuột và bàn phím WASD. Bảng 4.1 tóm tắt kết quả kiểm thử chức năng khởi động và điều khiển.")

append_table(
    "Bảng 4.1. Kết quả kiểm thử khởi động và điều khiển",
    ["Mã TC", "Chức năng kiểm thử", "Kết quả thực tế", "Đánh giá"],
    [
        ["TC01", "Khởi động và nhận diện chế độ VR trên kính Meta Quest", "Kích hoạt đúng XR Origin, tắt camera PC, nhận diện tracking 6DoF của đầu và tay cầm", "Đạt"],
        ["TC02", "Khởi động chế độ màn hình phẳng trên máy tính PC", "Kích hoạt FPSController, điều khiển góc nhìn bằng chuột và di chuyển bằng phím WASD", "Đạt"],
        ["TC03", "Nhận tín hiệu các nút bấm tay cầm (Trigger, Grip, A/B/X/Y)", "Đọc đúng trạng thái nhấn nút qua New Input System và OpenXR", "Đạt"]
    ],
    [1.0, 3.2, 2.3, 0.7]
)

append_p(h3, "4.2.2. Kiểm thử di chuyển và tương tác VR")
append_p(body, "Kiểm thử di chuyển xác nhận người chơi có thể bước đi liên tục bằng cần gạt trái theo hướng nhìn của đầu với tốc độ 2.5 m/s, không bị xuyên qua tường hay bệ hiện vật nhờ hệ thống va chạm vật lý. Tính năng xoay nhân tạo bằng cần gạt phải được vô hiệu hóa để người dùng quay đầu tự nhiên trong không gian thực. Cửa kính trượt AutomaticSlidingDoor tự động mở khi người chơi đến gần trong cự ly 2.8m và tự đóng lại sau khi đi qua, có cơ chế an toàn giữ cửa mở khi có vật cản ở giữa. Nhân vật NPC (NpcThamQuan) di chuyển tuần tự giữa các bệ hiện vật trên NavMesh và dừng quan sát đúng thiết kế.")

append_table(
    "Bảng 4.2. Kết quả kiểm thử di chuyển và tương tác môi trường",
    ["Mã TC", "Chức năng kiểm thử", "Kết quả thực tế", "Đánh giá"],
    [
        ["TC04", "Di chuyển bằng cần joystick trái theo hướng nhìn", "Di chuyển bám sàn ở tốc độ 2.5 m/s, hướng đi trùng với hướng quay đầu HMD", "Đạt"],
        ["TC05", "Va chạm vật lý với tường và vật thể kiến trúc", "Người chơi bị cản lại chính xác khi chạm tường, cột, bệ, không lọt khỏi map", "Đạt"],
        ["TC06", "Cửa kính tự động trượt mở và cơ chế chống kẹp", "Cửa trượt êm khi người đến gần trong 2.8m, không đóng khi người đứng giữa", "Đạt"],
        ["TC07", "NPC tham quan di chuyển theo NavMesh", "NPC tự động đi qua các điểm trưng bày, dừng ngắm từ 4 đến 8 giây rồi đi tiếp", "Đạt"]
    ],
    [1.0, 3.2, 2.3, 0.7]
)

append_p(h3, "4.2.3. Kiểm thử Spatial UI và các chức năng chính")
append_p(body, "Giao diện không gian tại mỗi bệ trưng bày hiện vật (CubeExhibitionTrigger) hiển thị biểu tượng ngón tay chỉ dẫn nhấp nhô và luôn xoay hướng về phía người chơi. Khi dùng tia laser tay cầm bóp cò (hoặc nhấn phím Shift/E trên PC), bảng pop-up thuyết minh (MuseumPopupManager) mở ra giữa tầm nhìn trong thời gian 0.22 giây. Kiểm thử xác nhận toàn bộ 12 bệ trưng bày hiển thị chính xác hình ảnh và tiêu đề của 12 dân tộc Tây Nguyên (Ê Đê, Mạ, Xtiêng, Ba Na, Chư-Ru, M'Nông, Gia Rai, Brâu, Cơ Ho, Xơ Đăng, Giẻ Triêng, Rơ Măm). Pop-up tự đóng khi người chơi bấm nút đóng hoặc lùi ra xa bệ hơn 4.2m.")

append_table(
    "Bảng 4.3. Kết quả kiểm thử Spatial UI và bảng thông tin hiện vật",
    ["Mã TC", "Chức năng kiểm thử", "Kết quả thực tế", "Đánh giá"],
    [
        ["TC08", "Biểu tượng ngón tay chỉ dẫn tương tác trên bệ Cube", "Icon ngón tay xuất hiện trong cự ly 2.8m, nhấp nhô và luôn hướng về camera", "Đạt"],
        ["TC09", "Trỏ tia laser mở bảng ảnh thuyết minh pop-up", "Bóp cò tia laser mở pop-up giữa tầm nhìn mượt mà, nội dung ảnh sắc nét", "Đạt"],
        ["TC10", "Tính đầy đủ và chính xác thông tin 12 dân tộc", "Hiển thị đúng ảnh chi tiết và tiêu đề của toàn bộ 12 dân tộc Tây Nguyên", "Đạt"],
        ["TC11", "Đóng pop-up bằng nút bấm hoặc khoảng cách", "Bấm lại phím/cò hoặc lùi ra xa quá 4.2m pop-up tự thu lại, có bộ lọc debounce", "Đạt"]
    ],
    [1.0, 3.2, 2.3, 0.7]
)

append_p(h3, "4.2.4. Kiểm thử âm thanh và phản hồi xúc giác")
append_p(body, "Hệ thống âm thanh không gian 3D hoạt động dựa trên AudioListener gắn tại camera người chơi. Âm thanh video tư liệu cồng chiêng từ màn hình Tivi (TvVideoController) phát rõ khi ở gần và giảm dần theo khoảng cách, đồng thời đèn Ambilight tự đổi màu và độ sáng theo khung hình video. Phản hồi xúc giác qua VRInputHelper ghi nhận hai mức độ rung: rung nhẹ khi tia laser lướt qua bệ hiện vật (hover) và rung dứt khoát khi bóp cò chọn (select).")

# 4.3. Kiểm thử hiệu năng
append_p(h2, "4.3. Kiểm thử hiệu năng")

append_p(h3, "4.3.1. Kiểm thử FPS và Frame Time")
append_p(body, "Hiệu năng khung hình được đo đạc tại các khu vực khác nhau trong bảo tàng: khu sảnh đón tiếp, khu vực trưng bày 12 bệ hiện vật trung tâm, phòng chiếu Tivi và góc nhìn bao quát toàn cảnh. Công cụ tối ưu MuseumPerformanceOptimizer đã hỗ trợ tắt đổ bóng trên các đèn phụ, chuyển đổi các vật liệu cũ sang URP Lit và nướng Occlusion Culling để giảm số lượng Draw Calls. [Ghi chú: Nhóm bổ sung bảng số liệu FPS và Frame Time đo thực tế bằng Unity Profiler hoặc OVR Metrics Tool vào khung Bảng 4.4 dưới đây].")

append_table(
    "Bảng 4.4. Khung đo kiểm FPS và Frame Time tại các phân khu trong bảo tàng",
    ["Phân khu kiểm thử", "FPS trên kính VR", "Frame Time (ms)", "FPS trên PC", "Số Batches (Draw Calls)", "Ghi chú"],
    [
        ["1. Sảnh đón tiếp chính", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "Khu vực xuất phát ban đầu"],
        ["2. Khu trưng bày 12 bệ hiện vật", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "Nơi tập trung mô hình và bệ Cube"],
        ["3. Phòng chiếu Tivi tư liệu", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "Phát video và ánh sáng Ambilight"],
        ["4. Góc nhìn bao quát toàn bảo tàng", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "Tầm nhìn rộng nhiều vật thể"],
        ["5. Trạng thái hiển thị Pop-up chi tiết", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "[Chờ đo]", "Đang mở ảnh thuyết minh 2D"]
    ],
    [1.8, 1.1, 1.1, 1.0, 1.2, 1.0]
)

append_p(h3, "4.3.2. Kiểm thử độ trễ và khả năng phản hồi")
append_p(body, "Độ trễ chuyển động từ đầu đến màn hình (Motion-to-Photon) và độ trễ phản hồi thao tác bấm nút là yếu tố then chốt giúp trải nghiệm VR không bị giật lag. Ứng dụng tận dụng đường ống Universal Render Pipeline với chế độ Single-Pass Instanced Rendering để tối ưu thời gian dựng hình cho hai mắt. Trong các lần thử nghiệm thực tế, thao tác trỏ tia laser và mở bảng thông tin phản hồi gần như tức thì. [Ghi chú: Nhóm bổ sung số liệu đo độ trễ ms cụ thể nếu có công cụ đo chuyên dụng].")

append_p(h3, "4.3.3. Kiểm thử CPU, GPU và bộ nhớ")
append_p(body, "Quá trình kiểm thử theo dõi mức độ chiếm dụng phần cứng nhằm bảo đảm ứng dụng không làm quá nhiệt thiết bị khi chạy lâu. Mã nguồn các script tương tác chính được thiết kế để không tạo cấp phát bộ nhớ động trong vòng lặp Update, giúp hạn chế phát sinh thu gom rác (Garbage Collection spikes). Kích thước gói cài đặt APK standalone hiện tại là khoảng 2.05 GB. [Ghi chú: Nhóm bổ sung số liệu phần trăm CPU, GPU và dung lượng RAM đo thực tế qua Android Logcat hoặc Unity Profiler tại đây].")

# 4.4. Đánh giá trải nghiệm người dùng
append_p(h2, "4.4. Đánh giá trải nghiệm người dùng")

append_p(h3, "4.4.1. Đánh giá tính dễ sử dụng")
append_p(body, "Tính dễ sử dụng được đánh giá thông qua việc quan sát và lấy ý kiến của người tham gia trải nghiệm thử nghiệm ứng dụng bảo tàng. Các nội dung khảo sát tập trung vào: khả năng làm quen với cần gạt di chuyển, độ dễ dàng khi trỏ tia laser vào bệ hiện vật và thao tác đọc/đóng bảng thông tin thuyết minh. [Ghi chú: Nhóm bổ sung số lượng người tham gia thử nghiệm thực tế, kết quả khảo sát và điểm đánh giá theo thang đo khả dụng (ví dụ thang SUS) tại đây].")

append_p(h3, "4.4.2. Đánh giá mức độ tương tác và trực quan")
append_p(body, "Về khía cạnh trực quan, người dùng được hỏi về mức độ rõ nét của hình ảnh trang phục dân tộc, tính hữu ích của biểu tượng ngón tay 3D hướng dẫn và cảm nhận về không gian bảo tàng khi có thêm NPC và Tivi chiếu phim tư liệu. Phản hồi ban đầu cho thấy biểu tượng ngón tay giúp người tham quan dễ dàng nhận biết vị trí có thể tương tác mà không cần biển hướng dẫn bằng chữ. [Ghi chú: Nhóm bổ sung các nhận xét định tính hoặc tỷ lệ hài lòng thu thập được từ người dùng].")

append_p(h3, "4.4.3. Đánh giá mức độ thoải mái khi sử dụng VR")
append_p(body, "Đánh giá mức độ thoải mái tập trung vào việc theo dõi các triệu chứng mỏi mắt hoặc chóng mặt do chuyển động trong môi trường ảo. Việc loại bỏ tính năng xoay nhân tạo bằng cần gạt (chỉ giữ xoay đầu tự nhiên) và giữ tốc độ bước đi ở mức vừa phải (2.5 m/s) là các yếu tố quan trọng giúp giảm thiểu cảm giác mất thăng bằng. [Ghi chú: Nhóm bổ sung ghi nhận thực tế về tỷ lệ người dùng cảm thấy thoải mái sau phiên trải nghiệm 10-15 phút tại đây].")

# 4.5. Tổng hợp và đánh giá kết quả
append_p(h2, "4.5. Tổng hợp và đánh giá kết quả")

append_p(h3, "4.5.1. Tổng hợp kết quả kiểm thử")
append_p(body, "Tổng hợp quá trình kiểm thử cho thấy ứng dụng đã đáp ứng các yêu cầu cơ bản về mặt chức năng và khả năng vận hành ổn định trên cả hai môi trường VR và PC. Các thành phần tương tác chính gồm di chuyển, trỏ raycast, hiển thị bảng thuyết minh 12 dân tộc, cửa tự động và phát video đều hoạt động đúng thiết kế. Bảng 4.5 tóm tắt tình trạng nghiệm thu các hạng mục kiểm thử.")

append_table(
    "Bảng 4.5. Tổng hợp tình trạng nghiệm thu các hạng mục kiểm thử",
    ["Hạng mục kiểm thử", "Nội dung đánh giá", "Tình trạng", "Ghi chú"],
    [
        ["Khởi động và điều khiển", "Nhận diện chế độ VR/PC, điều khiển di chuyển và góc nhìn", "Đạt", "Vận hành ổn định"],
        ["Tương tác bệ hiện vật", "Tia laser, icon chỉ dẫn và mở pop-up 12 dân tộc", "Đạt", "Đầy đủ dữ liệu"],
        ["Môi trường phụ trợ", "Cửa trượt kính chống kẹp, NPC tham quan, Tivi tư liệu", "Đạt", "Hoạt động đúng kịch bản"],
        ["Hiệu năng khung hình", "Tốc độ khung hình và thời gian dựng hình trên thiết bị", "[Chờ bổ sung]", "Nhóm đo đạc số liệu chi tiết"],
        ["Trải nghiệm người dùng", "Mức độ dễ sử dụng, độ trực quan và cảm giác thoải mái", "[Chờ bổ sung]", "Nhóm tổng hợp từ khảo sát thực tế"]
    ],
    [1.8, 3.0, 1.1, 1.3]
)

append_p(h3, "4.5.2. Những kết quả đạt được")
append_p(body, "Đề tài đã đạt được các kết quả nổi bật gồm: (1) Xây dựng thành công không gian bảo tàng ảo trưng bày trang phục truyền thống 12 dân tộc Tây Nguyên với chất lượng đồ họa URP đồng bộ; (2) Hỗ trợ cơ chế vận hành kép linh hoạt giữa kính VR Meta Quest và máy tính cá nhân; (3) Hiện thực hóa các tương tác trực quan như tia laser kèm phản hồi rung haptic, pop-up thông tin tự co giãn, cửa tự động và NPC tăng tính sống động; (4) Chuẩn bị sẵn quy trình tối ưu hóa kết xuất giúp giảm thiểu Draw Calls phục vụ triển khai trên thiết bị di động.")

append_p(h3, "4.5.3. Hạn chế của ứng dụng")
append_p(body, "Bên cạnh các kết quả đạt được, ứng dụng hiện còn một số hạn chế: (1) Dung lượng gói cài đặt APK standalone còn tương đối lớn (khoảng 2.05 GB) do chứa nhiều tài nguyên hình ảnh chất lượng cao và video tư liệu chưa qua nén chuyên sâu; (2) Tương tác với hiện vật hiện dừng lại ở việc xem thông tin thuyết minh 2D, chưa hỗ trợ cầm nắm xoay mô hình 3D hay mặc thử ảo trên nhân vật; (3) Chưa hỗ trợ tính năng kết nối mạng đa người tham quan cùng lúc (Multiplayer); (4) Cử động của NPC tham quan còn ở mức mô phỏng chuyển động cơ bản, chưa áp dụng hệ thống animation xương hay Inverse Kinematics phức tạp.")

# ---------------------------------------------------------
# 3. Save Output Files
# ---------------------------------------------------------
os.makedirs(os.path.dirname(OUTPUT_QA), exist_ok=True)
doc.save(OUTPUT_QA)
print(f"Saved: {OUTPUT_QA}")

doc.save(OUTPUT_ROOT)
print(f"Saved: {OUTPUT_ROOT}")
print("Successfully built concise Chapter 4!")
