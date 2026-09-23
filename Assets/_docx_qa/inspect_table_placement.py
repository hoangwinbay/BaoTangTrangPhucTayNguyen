import docx
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls

doc = docx.Document(r'D:\Dự án bảo tàng\Bao_cao_Mo_phong_Bao_tang_Trang_phuc_Tay_Nguyen_Chuong_3.docx')

# Check how tables are placed in the document
for i, table in enumerate(doc.tables):
    parent = table._tbl.getparent()
    prev_sib = table._tbl.getprevious()
    next_sib = table._tbl.getnext()
    print(f"Table {i} parent tag: {parent.tag.split('}')[-1]}")
    if prev_sib is not None and prev_sib.tag.endswith('p'):
        p_prev = doc.paragraphs[[p._p for p in doc.paragraphs].index(prev_sib)]
        print(f"   Prev paragraph: {p_prev.text}")

print("Done inspecting table placement.")
