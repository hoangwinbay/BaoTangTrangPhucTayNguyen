import docx
from docx.oxml.ns import qn

DOC_PATH = r"D:\Dự án bảo tàng\Assets\_docx_qa\Bao_cao_Mo_phong_Bao_tang_Trang_phuc_Tay_Nguyen_Chuong_4.docx"

doc = docx.Document(DOC_PATH)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")
print(f"Total Sections: {len(doc.sections)}")

# 1. Verify TOC
print("\n=== TOC ENTRIES FOR CHAPTER 4 ===")
for i, p in enumerate(doc.paragraphs):
    if "CHƯƠNG 4" in p.text or ("4." in p.text and "\t" in p.text):
        print(f"[{i}] {repr(p.text)}")

# 2. Verify Chapter 4 Headings in Body
print("\n=== CHAPTER 4 HEADINGS IN BODY ===")
for i, p in enumerate(doc.paragraphs):
    if any(p.text.startswith(prefix) for prefix in ["CHƯƠNG 4", "4.1.", "4.2.", "4.3.", "4.4.", "4.5."]):
        first_run = p.runs[0] if p.runs else None
        sz = first_run.font.size.pt if first_run and first_run.font and first_run.font.size else "Default"
        bold = first_run.bold if first_run else None
        print(f"[{i}] (Size: {sz}, Bold: {bold}, Style: {p.style.name}): {p.text}")

# 3. Verify Tables in Chapter 4
print("\n=== TABLES IN CHAPTER 4 ===")
for i in range(6, len(doc.tables)):
    tbl = doc.tables[i]
    hdr = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
    r1 = [c.text.strip().replace('\n', ' ') for c in tbl.rows[1].cells]
    # Check shading of header
    tcPr = tbl.rows[0].cells[0]._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    shd_val = shd.get(qn('w:fill')) if shd is not None else None
    print(f"Table {i}: rows={len(tbl.rows)}, cols={len(tbl.columns)}, shd={shd_val}")
    print(f"   Header: {hdr[:3]}")
    print(f"   Row 1:  {r1[:3]}")

# 4. Check Body Paragraph formatting sample
print("\n=== SAMPLE BODY PARAGRAPH CHECK ===")
sample_p = None
for p in doc.paragraphs:
    if p.text.startswith("Mục tiêu trọng tâm của giai đoạn kiểm thử"):
        sample_p = p
        break

if sample_p:
    pPr = sample_p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    first_line = ind.get(qn('w:firstLine')) if ind is not None else None
    print(f"Body text sample: {sample_p.text[:60]}...")
    print(f"Style: {sample_p.style.name}, firstLine: {first_line}")
    for r in sample_p.runs[:2]:
        print(f"Run: font={r.font.name}, size={r.font.size.pt if r.font.size else 'default'}, bold={r.bold}")

print("\nVerification completed successfully!")
