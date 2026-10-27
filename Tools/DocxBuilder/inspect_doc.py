import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document(r'../../Docs/BaoCao/Chuong_3/Bao_cao_Mo_phong_Bao_tang_Trang_phuc_Tay_Nguyen_Chuong_3.docx')

print('=== SECTIONS INFO ===')
for i, s in enumerate(doc.sections):
    print(f'Section {i}: top={s.top_margin.cm:.2f}cm, bottom={s.bottom_margin.cm:.2f}cm, left={s.left_margin.cm:.2f}cm, right={s.right_margin.cm:.2f}cm, page_width={s.page_width.cm:.2f}cm, page_height={s.page_height.cm:.2f}cm')

print('\n=== CHECKING SAMPLES OF HEADING AND BODY ===')
samples = [
    ('CHƯƠNG 1', 92),
    ('CHƯƠNG 2', 127),
    ('CHƯƠNG 3', 187),
    ('Heading 2.1', 128),
    ('Heading 2.1.1', 129),
    ('Heading 3.1', 188),
    ('Heading 3.1.1', 189),
    ('Body Ch2', 130),
    ('Body Ch3', 190)
]

for label, idx in samples:
    p = doc.paragraphs[idx]
    r = p.runs[0] if p.runs else None
    rPr = r._r.get_or_add_rPr() if r is not None else None
    rFonts = rPr.find(qn('w:rFonts')) if rPr is not None else None
    font_ascii = rFonts.get(qn('w:ascii')) if rFonts is not None else (r.font.name if r else None)
    sz = rPr.find(qn('w:sz')) if rPr is not None else None
    sz_val = int(sz.get(qn('w:val')))/2 if sz is not None else (r.font.size.pt if r and r.font.size else None)
    b = rPr.find(qn('w:b')) is not None if rPr is not None else (r.bold if r else None)
    pPr = p._p.get_or_add_pPr()
    jc = pPr.find(qn('w:jc'))
    align = jc.get(qn('w:val')) if jc is not None else str(p.alignment)
    sp = pPr.find(qn('w:spacing'))
    spacing_line = sp.get(qn('w:line')) if sp is not None else None
    spacing_before = sp.get(qn('w:before')) if sp is not None else None
    spacing_after = sp.get(qn('w:after')) if sp is not None else None
    ind = pPr.find(qn('w:ind'))
    first_line = ind.get(qn('w:firstLine')) if ind is not None else None
    print(f'{label} [{idx}]: text="{p.text[:40]}"')
    print(f'   style={p.style.name}, font={font_ascii}, sz={sz_val}pt, bold={b}, align={align}')
    print(f'   spacing: line={spacing_line}, before={spacing_before}, after={spacing_after}, firstLine={first_line}')

print('\n=== TABLES INFO ===')
print('Number of tables:', len(doc.tables))
for i, table in enumerate(doc.tables):
    print(f'Table {i}: rows={len(table.rows)}, cols={len(table.columns)}')
    for r in range(min(2, len(table.rows))):
        row_txt = [cell.text.strip().replace('\n', ' ') for cell in table.rows[r].cells]
        print(f'   row {r}: {row_txt[:4]}')
