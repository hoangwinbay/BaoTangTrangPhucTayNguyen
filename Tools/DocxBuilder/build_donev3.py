# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Load REWRITTEN and TARGET_BLOCKS from build_donev2
import build_donev2

def run():
    src_path = r"C:\Users\MD15\Downloads\done.docx"
    dst_path = r"C:\Users\MD15\Downloads\donev3.docx"

    doc = docx.Document(src_path)
    print(f"Loaded {src_path}, total paragraphs: {len(doc.paragraphs)}")

    paragraphs_to_remove = []

    for key, p_indices in build_donev2.TARGET_BLOCKS:
        new_texts = build_donev2.REWRITTEN[key]
        assert len(new_texts) <= len(p_indices), f"Key {key} has more new texts than slots!"
        
        # Update existing paragraphs
        for idx, text in enumerate(new_texts):
            p_idx = p_indices[idx]
            p = doc.paragraphs[p_idx]
            p.text = text
            # Strict 1.5 line spacing matching paragraph 225
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = None
            p.paragraph_format.space_after = None
            
            # Ensure w:spacing has line="360" and lineRule="auto", and no w:after or w:before
            pPr = p._p.pPr
            if pPr.spacing is not None:
                # Remove after and before attributes if present
                for attr in ['after', 'before']:
                    key_qname = f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}"
                    if key_qname in pPr.spacing.attrib:
                        del pPr.spacing.attrib[key_qname]
            
            for r in p.runs:
                r.font.name = None
                r.font.size = None
                
        # Mark extra paragraphs for removal
        for extra_idx in p_indices[len(new_texts):]:
            paragraphs_to_remove.append(doc.paragraphs[extra_idx])

    # Remove extra paragraphs
    for p in paragraphs_to_remove:
        p_elem = p._p
        parent = p_elem.getparent()
        if parent is not None:
            parent.remove(p_elem)

    print(f"Successfully modified content and removed {len(paragraphs_to_remove)} redundant paragraphs.")
    print(f"Remaining paragraphs: {len(doc.paragraphs)}")

    doc.save(dst_path)
    print(f"Saved to {dst_path}")

if __name__ == "__main__":
    run()
