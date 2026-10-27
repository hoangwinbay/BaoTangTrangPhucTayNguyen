import docx
from docx.oxml.ns import qn

doc = docx.Document(r'../../Docs/BaoCao/Chuong_3/Bao_cao_Mo_phong_Bao_tang_Trang_phuc_Tay_Nguyen_Chuong_3.docx')

table = doc.tables[0]
print('Table 0 style:', table.style.name)
for r_idx in range(min(2, len(table.rows))):
    row = table.rows[r_idx]
    for c_idx, cell in enumerate(row.cells):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find(qn('w:shd'))
        shd_val = shd.get(qn('w:fill')) if shd is not None else None
        p = cell.paragraphs[0]
        font_name = p.runs[0].font.name if p.runs and p.runs[0].font else None
        sz = p.runs[0].font.size.pt if p.runs and p.runs[0].font and p.runs[0].font.size else None
        bold = p.runs[0].bold if p.runs else None
        txt = cell.text.strip().replace('\n', ' ')
        print(f'R{r_idx}C{c_idx}: text="{txt[:35]}", shd={shd_val}, font={font_name}, sz={sz}, bold={bold}, align={p.alignment}')
